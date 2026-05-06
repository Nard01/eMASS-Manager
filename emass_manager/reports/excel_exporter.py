from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

def export_table(rows:list[dict], destination:Path, sheet_name='Report'):
    wb=Workbook(); ws=wb.active; ws.title=sheet_name
    if not rows: rows=[{"message":"No data"}]
    headers=list(rows[0].keys()); ws.append(headers)
    for c in ws[1]: c.font=Font(bold=True)
    for r in rows: ws.append([r.get(h,'') for h in headers])
    ws.auto_filter.ref=ws.dimensions; ws.freeze_panes='A2'
    for col in ws.columns: ws.column_dimensions[col[0].column_letter].width=22
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            if isinstance(cell.value,str) and 'overdue' in cell.value.lower():
                cell.fill=PatternFill('solid',fgColor='7F1D1D')
    destination.parent.mkdir(parents=True,exist_ok=True); wb.save(destination)
